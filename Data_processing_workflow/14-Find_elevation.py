import os
import zipfile
from pathlib import Path
import openpyxl
import csv
from io import TextIOWrapper

def read_zipped_files(main_path):
    """
    Read zipped files in the main path, find BIF files, and extract LOCATION_ELEV values
    Using openpyxl directly instead of pandas
    """
    main_directory = Path(main_path)
    
    # Check if path exists
    if not main_directory.exists():
        print(f"Error: Path {main_path} does not exist!")
        return
    
    # Results storage
    results = []
    failed_sites = []
    no_elevation_sites = []
    no_bif_sites = []
    
    # Find all zip files in the main directory
    zip_files = [f for f in main_directory.glob("*.zip") if f.is_file()]
    
    if not zip_files:
        print(f"No zip files found in {main_path}")
        return
    
    print(f"Found {len(zip_files)} zip file(s) to process\n")
    print("="*80)
    
    for zip_path in zip_files:
        try:
            # Open the zip file
            with zipfile.ZipFile(zip_path, 'r') as zip_ref:
                # Get all files in the zip
                all_files = zip_ref.namelist()
                
                # Find files with 'BIF' in the name (case insensitive)
                bif_files = [f for f in all_files if 'BIF' in f.upper()]
                
                if not bif_files:
                    print(f"⚠️  No BIF file found in {zip_path.name}")
                    no_bif_sites.append(zip_path.name)
                    continue
                
                # Process each BIF file
                for bif_file in bif_files:
                    site_name = "Unknown"
                    elevation = None
                    
                    try:
                        # Check file extension
                        if bif_file.endswith('.xlsx'):
                            # Read Excel file using openpyxl
                            with zip_ref.open(bif_file) as excel_file:
                                # Load workbook from the file-like object
                                workbook = openpyxl.load_workbook(excel_file, data_only=True)
                                sheet = workbook.active
                                
                                # Find headers (assuming first row contains column names)
                                headers = []
                                for col in range(1, sheet.max_column + 1):
                                    cell_value = sheet.cell(1, col).value
                                    if cell_value:
                                        headers.append(str(cell_value).strip())
                                    else:
                                        headers.append(f"Column_{col}")
                                
                                # Find column indices for SITE_ID, VARIABLE, DATAVALUE
                                site_id_col = None
                                variable_col = None
                                datavalue_col = None
                                
                                for idx, header in enumerate(headers):
                                    if header == 'SITE_ID':
                                        site_id_col = idx + 1  # openpyxl is 1-indexed
                                    elif header == 'VARIABLE':
                                        variable_col = idx + 1
                                    elif header == 'DATAVALUE':
                                        datavalue_col = idx + 1
                                
                                if variable_col is None or datavalue_col is None:
                                    print(f"⚠️  Required columns not found in {zip_path.name} - {bif_file}")
                                    failed_sites.append(f"{zip_path.name} - Missing required columns")
                                    continue
                                
                                # Search for LOCATION_ELEV
                                for row in range(2, sheet.max_row + 1):
                                    variable = sheet.cell(row, variable_col).value
                                    if variable and str(variable).strip() == 'LOCATION_ELEV':
                                        elevation = sheet.cell(row, datavalue_col).value
                                        break
                                    
                                    # Also get site_id if available
                                    if site_id_col and site_id_col is not None:
                                        site_val = sheet.cell(row, site_id_col).value
                                        if site_val and site_name == "Unknown":
                                            site_name = str(site_val).strip()
                                
                                # If site_id not found in data rows, try header row
                                if site_name == "Unknown" and site_id_col:
                                    site_name = str(sheet.cell(1, site_id_col).value).strip()
                        
                        elif bif_file.endswith('.csv'):
                            # Read CSV file
                            with zip_ref.open(bif_file) as csv_file:
                                # Try different encodings
                                for encoding in ['utf-8', 'latin1', 'cp1252', 'iso-8859-1']:
                                    try:
                                        csv_file.seek(0)
                                        text_stream = TextIOWrapper(csv_file, encoding=encoding)
                                        reader = csv.DictReader(text_stream, delimiter='\t')
                                        
                                        for row in reader:
                                            if row.get('VARIABLE') == 'LOCATION_ELEV':
                                                elevation = row.get('DATAVALUE')
                                            if site_name == "Unknown" and row.get('SITE_ID'):
                                                site_name = row.get('SITE_ID')
                                        
                                        if elevation is not None:
                                            break
                                    except:
                                        continue
                        else:
                            print(f"⚠️  Unknown file format in {zip_path.name}: {bif_file}")
                            failed_sites.append(f"{zip_path.name} - {bif_file} (unknown format)")
                            continue
                        
                        if elevation is not None:
                            results.append({
                                'site_name': site_name,
                                'elevation': elevation,
                                'zip_file': zip_path.name
                            })
                            print(f"✓ Site: {site_name}")
                            print(f"  Zip file: {zip_path.name}")
                            print(f"  Elevation: {elevation} meters")
                            print(f"  BIF file: {bif_file}")
                            print("-"*40)
                        else:
                            print(f"⚠️  LOCATION_ELEV not found in {zip_path.name} (Site: {site_name})")
                            no_elevation_sites.append(f"{zip_path.name} (Site: {site_name})")
                    
                    except Exception as e:
                        print(f"❌ Error reading {bif_file} in {zip_path.name}: {str(e)}")
                        failed_sites.append(f"{zip_path.name} - {bif_file}")
        
        except Exception as e:
            print(f"❌ Error processing zip file {zip_path.name}: {str(e)}")
            failed_sites.append(zip_path.name)
    
    # Print summary
    print("\n" + "="*80)
    print("SUMMARY")
    print("="*80)
    print(f"Total zip files processed: {len(zip_files)}")
    print(f"Successfully extracted elevation: {len(results)}")
    print(f"Missing LOCATION_ELEV: {len(no_elevation_sites)}")
    print(f"No BIF file found: {len(no_bif_sites)}")
    print(f"Failed/Errors: {len(failed_sites)}")
    
    if results:
        print("\n" + "-"*40)
        print("DETAILED RESULTS:")
        print("-"*40)
        for result in results:
            print(f"Site: {result['site_name']:15} | Elevation: {result['elevation']:>6} meters | Zip: {result['zip_file']}")
    
    if no_elevation_sites:
        print("\n" + "-"*40)
        print("SITES WITH NO LOCATION_ELEV:")
        print("-"*40)
        for site in no_elevation_sites:
            print(f"  • {site}")
    
    if no_bif_sites:
        print("\n" + "-"*40)
        print("SITES WITH NO BIF FILE:")
        print("-"*40)
        for site in no_bif_sites:
            print(f"  • {site}")
    
    if failed_sites:
        print("\n" + "-"*40)
        print("FAILED/PROBLEM SITES (first 20 shown):")
        print("-"*40)
        for site in failed_sites[:20]:
            print(f"  • {site}")
        if len(failed_sites) > 20:
            print(f"  ... and {len(failed_sites) - 20} more")
    
    return results, failed_sites, no_elevation_sites, no_bif_sites

def main():
    # Set the path
    base_path = r"M:\Research\WUE_CUE\ameri_data"
    
    # Run the extraction
    results, failed, no_elev, no_bif = read_zipped_files(base_path)
    
    # Save results to a file
    if results:
        output_file = Path(base_path) / "elevation_results.txt"
        with open(output_file, 'w') as f:
            f.write("Site Name\tElevation (m)\tZip File\n")
            for result in results:
                f.write(f"{result['site_name']}\t{result['elevation']}\t{result['zip_file']}\n")
        print(f"\nResults saved to: {output_file}")

if __name__ == "__main__":
    main()